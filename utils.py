"""Shared utility functions: file metadata extraction, filename generation, password validation, audit logging, rate limiting."""

import collections
import functools
import os
import re
import time
from datetime import datetime

from flask import session, request, redirect, url_for, jsonify, g

from models import get_db, FAMILY_ABBR

# --- Rate limiting for login ---
_login_attempts = collections.defaultdict(list)
LOGIN_RATE_LIMIT = 5
LOGIN_RATE_WINDOW = 300


def _check_rate_limit(ip):
    now = time.time()
    _login_attempts[ip] = [t for t in _login_attempts[ip] if now - t < LOGIN_RATE_WINDOW]
    if len(_login_attempts[ip]) >= LOGIN_RATE_LIMIT:
        return False
    _login_attempts[ip].append(now)
    return True


def get_org_id():
    """Get current organization ID from session, default to 1."""
    return session.get('org_id', 1)


def validate_password(password):
    errors = []
    if len(password) < 16:
        errors.append("Password must be at least 16 characters")
    if not re.search(r'[A-Z]', password):
        errors.append("Must contain an uppercase letter")
    if not re.search(r'[a-z]', password):
        errors.append("Must contain a lowercase letter")
    if not re.search(r'[0-9]', password):
        errors.append("Must contain a number")
    if not re.search(r'[^A-Za-z0-9]', password):
        errors.append("Must contain a special character")
    return errors


def log_audit(action, target_type=None, target_id=None, details=None, conn=None):
    close = False
    if conn is None:
        conn = get_db()
        close = True
    conn.execute(
        "INSERT INTO audit_log (user_id, username, action, target_type, target_id, details, timestamp) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (session.get('user_id'), session.get('username'), action, target_type,
         str(target_id) if target_id is not None else None, details,
         datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    )
    conn.commit()
    if close:
        conn.close()


def _extract_file_created(filepath, ext):
    """Extract creation/modification date from file metadata."""
    try:
        if ext in ('.jpg', '.jpeg', '.png', '.gif', '.webp'):
            from PIL import Image
            from PIL.ExifTags import Base as ExifBase
            img = Image.open(filepath)
            exif = img.getexif()
            if exif:
                for tag in (ExifBase.DateTimeOriginal, ExifBase.DateTimeDigitized, ExifBase.DateTime):
                    val = exif.get(tag)
                    if val:
                        return val.replace(":", "-", 2)
            img.close()

        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(filepath)
            info = reader.metadata
            if info:
                for field in (info.creation_date, info.modification_date):
                    if field:
                        return field.strftime("%Y-%m-%d %H:%M:%S")

        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            props = doc.core_properties
            if props.created:
                return props.created.strftime("%Y-%m-%d %H:%M:%S")
            if props.modified:
                return props.modified.strftime("%Y-%m-%d %H:%M:%S")

        elif ext == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(filepath, read_only=True)
            props = wb.properties
            if props.created:
                return props.created.strftime("%Y-%m-%d %H:%M:%S")
            if props.modified:
                return props.modified.strftime("%Y-%m-%d %H:%M:%S")
            wb.close()

        elif ext == '.pptx':
            import zipfile
            from xml.etree import ElementTree
            with zipfile.ZipFile(filepath) as z:
                if 'docProps/core.xml' in z.namelist():
                    tree = ElementTree.parse(z.open('docProps/core.xml'))
                    ns = {'dcterms': 'http://purl.org/dc/terms/'}
                    created = tree.find('.//dcterms:created', ns)
                    if created is not None and created.text:
                        return created.text[:19].replace("T", " ")

    except Exception:
        pass
    return None


def admin_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        user_id = session.get('user_id')
        if not user_id:
            return redirect(url_for('auth.login'))
        conn = get_db()
        user = conn.execute("SELECT role FROM users WHERE id = ?", (user_id,)).fetchone()
        conn.close()
        if not user or user['role'] != 'admin':
            if request.path.startswith('/api/'):
                return jsonify({"error": "Admin access required"}), 403
            return redirect(url_for('dashboard.dashboard'))
        return f(*args, **kwargs)
    return decorated


def _generate_artifact_filename(conn, objective_id, domain_name, ext):
    """Generate auto-renamed filename: AC-3.01.01.a-Domain{ext}"""
    clean_id = objective_id.replace("[", ".").replace("]", "").replace(" ", "").replace("\t", "")
    obj_row = conn.execute("SELECT family FROM objectives WHERE id = ?", (objective_id,)).fetchone()
    abbr = FAMILY_ABBR.get(obj_row["family"], "") if obj_row else ""
    prefix = f"{abbr}-{clean_id}" if abbr else clean_id

    domain_part = ""
    if domain_name:
        domain_part = "-" + domain_name.replace(" ", "-")

    return f"{prefix}{domain_part}{ext}"
